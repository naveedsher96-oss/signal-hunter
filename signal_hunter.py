"""
Signal Hunter - warm, intent-based leads for remote Sales, CRM & VA work
------------------------------------------------------------------------
Finds companies (and people) who have PUBLICLY shown they need help with
GoHighLevel, CRM automation, sales / business development, customer success,
virtual / admin assistance or data entry - and drafts an email that answers
that need.

Signals it collects (free, public sources only):
  * Asking for help    - Reddit posts ([Hiring], "looking for a GHL expert"...),
                         Hacker News comments and "Who is hiring" threads
  * Job posts          - Greenhouse, Lever, Ashby, Workable, SmartRecruiters,
                         Recruitee boards + RemoteOK, Remotive, We Work Remotely,
                         Himalayas remote feeds
  * Tools they use     - GoHighLevel / HubSpot / Zoho / Salesforce... on their site
  * News               - hiring news, funding, new sales leaders, CRM rollouts
Each signal is weighted by how recent it is.

Two modes:
  python signal_hunter.py my_list.xlsx     scan YOUR company list
  python signal_hunter.py --discover       find warm companies with no list

Most people use the app instead: double-click "Start Signal Hunter.bat".
Nothing is ever sent. You review every email before sending it yourself.
"""

import csv
import html
import json
import os
import re
import sys
import time
import urllib.parse
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta, timezone
from email.utils import parsedate_to_datetime

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(BASE_DIR, "output")
EMAIL_DIR = os.path.join(OUT_DIR, "emails")
SEEN_FILE = os.path.join(OUT_DIR, "seen_signals.json")
HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) SignalHunter/2.0 "
                         "(personal job research)"}
TIMEOUT = 15
SESSION = requests.Session()

# ------------------------------------------------------------- defaults --

DEFAULT_CATEGORIES = {
    "GoHighLevel": ["gohighlevel", "go high level", "go highlevel", "highlevel", "ghl",
                    "leadconnector", "lead connector"],
    "CRM Automation": ["crm", "hubspot", "zoho", "salesforce", "pipedrive", "zapier",
                       "make.com", "n8n", "crm automation", "sales automation",
                       "business automation", "workflow automation", "revops", "revenue operations", "sales operations",
                       "sales ops", "marketing automation", "activecampaign", "keap"],
    "Sales & Business Development": ["sales", "business development", "sdr", "bdr",
                                     "account executive", "appointment setter",
                                     "appointment setting", "cold caller", "cold calling",
                                     "lead generation", "lead gen", "closer",
                                     "inside sales", "sales rep"],
    "Customer Success": ["customer success", "customer support", "customer service",
                         "account manager", "client success", "onboarding specialist",
                         "support specialist"],
    "Virtual / Admin Assistant": ["virtual assistant", "va", "admin assistant",
                                  "administrative assistant", "executive assistant",
                                  "data entry", "operations assistant",
                                  "personal assistant", "admin support",
                                  "back office"],
}
DEFAULT_SUBREDDITS = ["forhire", "hiring", "gohighlevel", "VirtualAssistant",
                      "smallbusiness", "Entrepreneur", "EntrepreneurRideAlong",
                      "startups", "sweatystartup", "Solopreneur", "ecommerce", "shopify",
                      "agency", "realtors", "coaching", "sales", "hubspot", "zoho",
                      "salesforce", "marketing", "slavelabour"]

# Owners describing a STRUGGLE (no "hiring" word needed) -> which service fixes it
DEFAULT_PAIN = {
    "Virtual / Admin Assistant": [
        "overwhelmed", "drowning in", "too much admin", "buried in emails",
        "can't keep up with emails", "cant keep up with emails", "no time for admin",
        "wearing all the hats", "wearing every hat", "doing everything myself",
        "do everything myself", "burnt out", "burned out", "burnout", "paperwork",
        "inbox is a mess", "need to delegate", "how do i delegate", "first hire",
        "can't afford to hire", "cant afford to hire", "spending hours on admin"],
    "Sales & Business Development": [
        "no clients", "no customers", "no sales", "need more clients",
        "need more customers", "need more leads", "get more clients",
        "get more customers", "find clients", "find customers", "first client",
        "first clients", "first customers", "sales are slow", "sales have dropped",
        "sales are down", "slow month", "struggling to get", "struggling to find",
        "struggling to close", "how do i get clients", "how do i get customers",
        "how to get clients", "how to get customers", "no leads", "pipeline is empty",
        "cold outreach isn't working", "not getting any leads"],
    "CRM Automation": [
        "falling through the cracks", "fall through the cracks", "slipping through",
        "losing leads", "lost leads", "forget to follow up", "forgot to follow up",
        "follow ups are a mess", "follow-ups are a mess", "track leads in a spreadsheet",
        "tracking leads in excel", "tracking customers in a spreadsheet",
        "spreadsheet to track", "messy pipeline", "which crm", "what crm",
        "best crm for", "crm recommendation", "automate my", "manual work",
        "repetitive tasks"],
    "Customer Success": [
        "customers are complaining", "losing customers", "customer churn",
        "customers keep leaving", "can't respond fast enough", "support tickets piling",
        "bad reviews", "refund requests", "customer retention"],
}
DEFAULT_SMALLBIZ = [
    "small business", "my business", "my small business", "my company", "my startup",
    "just started", "just launched", "recently started", "recently launched",
    "new business", "started my", "launched my", "first year", "startup", "start-up",
    "solopreneur", "solo founder", "founder", "co-founder", "one-man", "one man band",
    "small team", "side hustle", "bootstrapped", "local business", "family business",
    "my shop", "my store", "my agency", "my practice", "my salon", "my clinic",
    "my restaurant", "my brand", "self-employed", "self employed", "owner"]
SMALLBIZ_SUBS = {"smallbusiness", "entrepreneur", "entrepreneurridealong", "startups",
                 "sweatystartup", "solopreneur", "ecommerce", "shopify", "realtors",
                 "coaching", "agency"}
SMALL_SITE_PATTERNS = {"Wix": ["wix.com", "wixstatic", "_wixcss"],
                       "Squarespace": ["squarespace.com", "static1.squarespace"],
                       "GoDaddy Website Builder": ["img1.wsimg.com", "godaddy"],
                       "Weebly": ["weebly.com", "editmysite"],
                       "Carrd": ["carrd.co"], "Shopify": ["cdn.shopify.com"],
                       "Linktree": ["linktr.ee"]}
SEG_SMALL = "Small business / founder"
DEFAULT_HN_KEYWORDS = ["gohighlevel", "virtual assistant", "crm", "revops",
                       "sales operations", "customer success", "data entry", "hubspot",
                       "salesforce admin", "zoho", "business development"]

# Someone asking for help / hiring
INTENT_RE = re.compile(
    r"(\[hiring\]|\bhiring\b|looking for|look(ing)? to hire|need (a|an|some|someone|help)|"
    r"\bneeded\b|seeking|searching for|in search of|want(ing)? to hire|recommend(ation)?s?|"
    r"anyone (know|have|recommend)|who can (help|build|set)|can someone|help me (with|set)|"
    r"outsourc|freelancer (needed|wanted)|\bwanted\b|paying)", re.I)
# People OFFERING services (competitors) - skip these
OFFER_RE = re.compile(r"(\[for hire\]|\bfor hire\b|\[offer\]|i('| a)m a (freelance|va|"
                      r"virtual assistant)|i offer|my services|hire me|available for work)",
                      re.I)

# ------------------------------------------------------------- remote rules --

WORLDWIDE = re.compile(
    r"(worldwide|anywhere|global|globally|international|apac|asia|south asia|"
    r"pakistan|islamabad|lahore|karachi|middle east|mena|gcc|uae|any location|"
    r"all locations|remote\s*-\s*remote|distributed|world)", re.I)
MAYBE = re.compile(r"\b(emea|europe|eu|cet|gmt|utc)\b", re.I)
REGION_LOCK = re.compile(
    r"\b(us|usa|u\.s\.|united states|canada|uk|united kingdom|england|ireland|"
    r"germany|france|spain|netherlands|australia|new zealand|latam|latin america|"
    r"brazil|mexico|india|philippines|north america|americas|"
    r"new york|san francisco|london|berlin|toronto|austin|boston|chicago|seattle)\b",
    re.I)
REMOTE_WORD = re.compile(r"remote|work from home|wfh|distributed|anywhere|telecommut", re.I)

# ------------------------------------------------------------- scoring --

KIND_POINTS = {
    "ask": 50,            # publicly asking for help / hiring post
    "job_eligible": 50,   # remote job, Pakistan-eligible
    "job_remote": 30,     # remote job, location unclear
    "struggle": 40,       # owner describing a pain we solve
    "hiring_news": 20,
    "job_locked": 20,     # remote but US/UK only -> pitch as contractor
    "tech_ghl": 20,
    "crm_news": 15,
    "funding": 15,
    "tech": 12,
    "mention": 12,        # talking about the topic
    "careers": 12,        # careers page mentions the need
    "job_onsite": 10,     # office job -> pitch remote contractor
    "leader": 10,
    "expansion": 5,
    "small_site": 5,
}
SMALLBIZ_BONUS = 10
KIND_LABEL = {
    "ask": "Asking for help",
    "struggle": "Struggling with it",
    "small_site": "Small-business website",
    "job_eligible": "Remote job (Pakistan-eligible)",
    "job_remote": "Remote job (verify location)",
    "job_locked": "Remote job, region-locked",
    "job_onsite": "Job open, not remote",
    "hiring_news": "Hiring news",
    "tech_ghl": "Runs GoHighLevel",
    "tech": "Uses CRM / sales tool",
    "mention": "Talking about it",
    "careers": "Careers page mentions it",
    "crm_news": "CRM / tool adoption",
    "funding": "Funding raised",
    "leader": "New sales/revenue leader",
    "expansion": "Expansion / launch",
}
CATEGORY_CAP = 60

NEWS_RULES = [
    ("funding", re.compile(r"(raises|raised|secures|closes|lands|funding|series [a-e]\b|"
                           r"seed round|investment led|growth round)", re.I)),
    ("leader", re.compile(r"(appoints|names|hires|welcomes|joins as|promoted to).{0,60}"
                          r"(cro|chief revenue|vp of sales|vp sales|head of sales|"
                          r"chief sales|revenue operations|chief commercial|chief growth|"
                          r"customer success|vp of revenue)", re.I)),
    ("hiring_news", re.compile(r"(hiring|to hire|jobs|recruit|headcount|job openings)",
                               re.I)),
    ("crm_news", re.compile(r"(hubspot|salesforce|zoho|gohighlevel|highlevel|pipedrive|"
                            r"dynamics 365|crm (migration|implementation|rollout)|"
                            r"selects .* crm)", re.I)),
    ("expansion", re.compile(r"(expands|expansion|launches|new office|enters .* market|"
                             r"acquires|acquisition|opens)", re.I)),
]

TECH_PATTERNS = {
    "GoHighLevel": ["leadconnectorhq", "msgsndr.com", "gohighlevel", "highlevel.com",
                    "link.msgsndr", "api.leadconnector"],
    "HubSpot": ["js.hs-scripts.com", "hs-analytics", "hsforms", "js.hubspot.com",
                "hs-banner"],
    "Zoho": ["salesiq.zoho", "zohopublic", "zcrm", "zohocdn", "crm.zoho"],
    "Salesforce": ["pardot", "force.com", "salesforce.com/", "salesforceliveagent"],
    "Pipedrive": ["pipedrive"],
    "ActiveCampaign": ["activecampaign", "trackcmp.net"],
    "Keap": ["infusionsoft", "keap.app"],
    "Calendly": ["assets.calendly.com"],
}

# ------------------------------------------------------------------ helpers --

def http_get(url, as_json=True, headers=None):
    """Single network entry point (patched in tests)."""
    try:
        r = SESSION.get(url, headers={**HEADERS, **(headers or {})}, timeout=TIMEOUT)
        if r.status_code != 200:
            return None
        return r.json() if as_json else r.text
    except (requests.RequestException, ValueError):
        return None


def strip_html(s):
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", s or ""))).strip()


def to_date(v):
    """Anything date-like -> 'YYYY-MM-DD' (or '')."""
    if v in (None, ""):
        return ""
    try:
        if isinstance(v, (int, float)) or str(v).isdigit():
            n = float(v)
            n = n / 1000 if n > 1e11 else n
            return datetime.fromtimestamp(n, timezone.utc).strftime("%Y-%m-%d")
        s = str(v)
        if re.match(r"\d{4}-\d{2}-\d{2}", s):
            return s[:10]
        return parsedate_to_datetime(s).strftime("%Y-%m-%d")
    except (TypeError, ValueError, OverflowError, OSError):
        return ""


def days_ago(date):
    try:
        return (datetime.now().date() - datetime.strptime(date, "%Y-%m-%d").date()).days
    except (TypeError, ValueError):
        return None


def recency(date):
    d = days_ago(date)
    if d is None:
        return 0.6
    return 1.0 if d <= 14 else 0.8 if d <= 30 else 0.55 if d <= 60 else 0.3


def in_window(date, lookback):
    d = days_ago(date)
    return d is None or d <= lookback


SUFFIXES = (" inc", " ltd", " llc", " limited", " corp", " co", " gmbh", " technologies",
            " technology", " labs", " hq", " io", " ai", " group", " software")


def norm(name):
    n = re.sub(r"[^a-z0-9 ]", " ", (name or "").lower())
    n = re.sub(r"\s+", " ", n).strip()
    changed = True
    while changed:
        changed = False
        for s in SUFFIXES:
            if n.endswith(s) and len(n) > len(s) + 1:
                n, changed = n[: -len(s)].strip(), True
    return n


def same_company(a, b):
    a, b = norm(a), norm(b)
    if not a or not b:
        return False
    return a == b or (min(len(a), len(b)) >= 4 and (a.startswith(b + " ") or
                                                     b.startswith(a + " ")))


def name_regex(company):
    return re.compile(r"(?<![a-z0-9])" + re.escape(norm(company)).replace(r"\ ", r"\s+")
                      + r"(?![a-z0-9])", re.I)


def compile_categories(cfg):
    cats = cfg.get("categories") or DEFAULT_CATEGORIES
    out = {}
    for name, kws in cats.items():
        kws = [k.strip() for k in (kws if isinstance(kws, list) else str(kws).split(","))
               if k.strip()]
        if kws:
            out[name] = re.compile(r"(?<![a-z0-9])(" + "|".join(
                re.escape(k).replace(r"\ ", r"\s+") for k in kws) + r")(?![a-z0-9])", re.I)
    return out


def categorize(text, cats):
    """Matching categories, ordered by where they first appear in the text."""
    hits = [(m.start(), n) for n, rx in cats.items() for m in [rx.search(text or "")] if m]
    return [n for _, n in sorted(hits)]


def ev(kind, category, source, title, url, date, detail="", **extra):
    e = {"kind": kind, "category": category or "Growth", "source": source,
         "title": (title or "")[:200], "detail": (detail or title or "")[:300],
         "url": url or "", "date": date or ""}
    e.update(extra)
    return e

# ------------------------------------------------------------------ config --

def with_defaults(cfg):
    cfg = dict(cfg)
    cfg.setdefault("categories", DEFAULT_CATEGORIES)
    cfg.setdefault("subreddits", DEFAULT_SUBREDDITS)
    cfg.setdefault("hn_keywords", DEFAULT_HN_KEYWORDS)
    cfg.setdefault("lookback_days", 45)
    cfg.setdefault("discover_limit", 75)
    cfg.setdefault("pain_signals", DEFAULT_PAIN)
    cfg.setdefault("smallbiz_words", DEFAULT_SMALLBIZ)
    cfg.setdefault("delay_seconds", 1)
    return cfg


def load_config():
    path = os.path.join(BASE_DIR, "config.json")
    if not os.path.exists(path):
        path = os.path.join(BASE_DIR, "defaults.json")
    with open(path, encoding="utf-8") as f:
        return with_defaults(json.load(f))

# ------------------------------------------------------------ company lists --

def read_table(path):
    """Return (headers, rows) from a CSV or Excel file. Rows are dicts keyed by header."""
    raw = []
    if path.lower().endswith((".xlsx", ".xlsm")):
        ws = load_workbook(path, read_only=True, data_only=True).active
        raw = [["" if v is None else str(v).strip() for v in r]
               for r in ws.iter_rows(values_only=True)]
    else:
        for enc in ("utf-8-sig", "cp1252"):
            try:
                with open(path, newline="", encoding=enc) as f:
                    raw = [[(v or "").strip() for v in r] for r in csv.reader(f)]
                break
            except UnicodeDecodeError:
                continue
    raw = [r for r in raw if any(r)]
    if not raw:
        return [], []
    headers = [h or f"Column {i + 1}" for i, h in enumerate(raw[0])]
    rows = [{headers[i]: (r[i] if i < len(r) else "") for i in range(len(headers))}
            for r in raw[1:]]
    return headers, rows


COLUMN_HINTS = {
    "company": ["company", "company name", "organization", "organisation", "account",
                "account name", "employer", "business", "name"],
    "website": ["website", "domain", "url", "company website", "web", "site"],
    "ats_slug": ["ats_slug", "slug", "board slug"],
    "contact_name": ["contact_name", "contact", "contact name", "full name",
                     "first name", "person", "hiring manager"],
    "contact_email": ["contact_email", "email", "e-mail", "email address", "work email"],
}


def guess_columns(headers):
    low = {h.lower().strip(): h for h in headers}
    guess = {}
    for field, hints in COLUMN_HINTS.items():
        guess[field] = next((low[h] for h in hints if h in low), "")
        if not guess[field]:
            guess[field] = next((orig for l, orig in low.items()
                                 if any(h in l for h in hints[:3])), "")
    if not guess["company"] and headers:
        guess["company"] = headers[0]
    return guess


def to_companies(rows, mapping):
    out, seen = [], set()
    col = lambda r, f: (r.get(mapping.get(f) or "", "") or "").strip()
    for r in rows:
        name = col(r, "company")
        if not name or name.lower() in seen:
            continue
        seen.add(name.lower())
        out.append({"company": name, "website": col(r, "website"),
                    "ats_slug": col(r, "ats_slug"), "contact_name": col(r, "contact_name"),
                    "contact_email": col(r, "contact_email")})
    return out


def load_companies(path, mapping=None):
    headers, rows = read_table(path)
    return to_companies(rows, mapping or guess_columns(headers))

# ------------------------------------------------- company job boards (ATS) --

def domain_of(website):
    if not website:
        return ""
    host = urllib.parse.urlparse(website if "//" in website else "//" + website).netloc
    return host.lower().replace("www.", "")


def slug_candidates(company, website, explicit):
    if explicit:
        return [explicit.lower()]
    base = norm(company)
    c = [base.replace(" ", ""), base.replace(" ", "-")]
    d = domain_of(website)
    if d:
        c.append(d.split(".")[0])
    return list(dict.fromkeys(s for s in c if s))


def _job(title, location, workplace, url, posted):
    return {"title": title or "", "location": location or "", "workplace": workplace or "",
            "url": url or "", "posted": to_date(posted)}


def ats_greenhouse(slug):
    d = http_get(f"https://boards-api.greenhouse.io/v1/boards/{slug}/jobs")
    if not d or "jobs" not in d:
        return None
    return [_job(j.get("title"), (j.get("location") or {}).get("name"), "",
                 j.get("absolute_url"), j.get("updated_at")) for j in d["jobs"]]


def ats_lever(slug):
    d = http_get(f"https://api.lever.co/v0/postings/{slug}?mode=json")
    if not isinstance(d, list) or not d:
        return None
    out = []
    for j in d:
        cats = j.get("categories") or {}
        locs = cats.get("allLocations") or [cats.get("location", "")]
        out.append(_job(j.get("text"), ", ".join(l for l in locs if l),
                        j.get("workplaceType"), j.get("hostedUrl"), j.get("createdAt")))
    return out


def ats_ashby(slug):
    d = http_get(f"https://api.ashbyhq.com/posting-api/job-board/{slug}")
    if not d or "jobs" not in d:
        return None
    out = []
    for j in d["jobs"]:
        sec = [s.get("location", "") for s in j.get("secondaryLocations") or []]
        wp = j.get("workplaceType") or ("Remote" if j.get("isRemote") else "")
        out.append(_job(j.get("title"), ", ".join([j.get("location", "")] + sec), wp,
                        j.get("jobUrl"), j.get("publishedAt")))
    return out


def ats_workable(slug):
    d = http_get(f"https://apply.workable.com/api/v1/widget/accounts/{slug}")
    if not d or "jobs" not in d:
        return None
    return [_job(j.get("title"), ", ".join(x for x in (j.get("city"), j.get("country")) if x),
                 "Remote" if j.get("telecommuting") else "", j.get("url") or
                 f"https://apply.workable.com/{slug}/j/{j.get('shortcode', '')}",
                 j.get("published_on")) for j in d["jobs"]]


def ats_smartrecruiters(slug):
    d = http_get(f"https://api.smartrecruiters.com/v1/companies/{slug}/postings")
    if not d or not d.get("content"):
        return None
    out = []
    for j in d["content"]:
        loc = j.get("location") or {}
        out.append(_job(j.get("name"), ", ".join(x for x in (loc.get("city"),
                                                            loc.get("country")) if x),
                        "Remote" if loc.get("remote") else "",
                        f"https://jobs.smartrecruiters.com/{slug}/{j.get('id', '')}",
                        j.get("releasedDate")))
    return out


def ats_recruitee(slug):
    d = http_get(f"https://{slug}.recruitee.com/api/offers/")
    if not d or not d.get("offers"):
        return None
    return [_job(j.get("title"), j.get("location") or ", ".join(
        x for x in (j.get("city"), j.get("country")) if x),
                 "Remote" if j.get("remote") else "", j.get("careers_url"),
                 j.get("published_at")) for j in d["offers"]]


BOARDS = [("Greenhouse", ats_greenhouse), ("Lever", ats_lever), ("Ashby", ats_ashby),
          ("Workable", ats_workable), ("SmartRecruiters", ats_smartrecruiters),
          ("Recruitee", ats_recruitee)]


def find_jobs(company):
    for slug in slug_candidates(company["company"], company["website"], company["ats_slug"]):
        for board, fn in BOARDS:
            jobs = fn(slug)
            if jobs:
                return board, slug, jobs
    return None, None, []


def remote_status(job):
    text = f"{job['location']} {job['workplace']}"
    if not (REMOTE_WORD.search(text) or job["workplace"].lower() == "remote"):
        return "Not remote"
    if WORLDWIDE.search(text):
        return "Eligible"
    if REGION_LOCK.search(text):
        return "Region-locked"
    if MAYBE.search(text):
        return "Maybe (verify)"
    return "Likely (verify)"


JOB_KIND = {"Eligible": "job_eligible", "Likely (verify)": "job_remote",
            "Maybe (verify)": "job_remote", "Region-locked": "job_locked",
            "Not remote": "job_onsite"}


def job_events(jobs, cats, source, lookback):
    out = []
    for j in jobs:
        c = categorize(j["title"], cats)
        if not c or not in_window(j["posted"], max(lookback, 90)):
            continue
        st = remote_status(j)
        out.append(ev(JOB_KIND[st], c[0], source, j["title"], j["url"], j["posted"],
                      f"{j['title']} - {j['location'] or j['workplace'] or 'location n/a'} "
                      f"[{st}]", role=j["title"], remote=st))
    return out

# ------------------------------------------------------ remote job feeds --

def feed_remoteok():
    d = http_get("https://remoteok.com/api")
    if not isinstance(d, list):
        return []
    return [{"company": j.get("company", ""), **_job(j.get("position"),
             j.get("location") or "Worldwide", "Remote", j.get("url"), j.get("date")),
             "source": "RemoteOK"} for j in d if isinstance(j, dict) and j.get("position")]


def feed_remotive():
    d = http_get("https://remotive.com/api/remote-jobs")
    if not d or "jobs" not in d:
        return []
    return [{"company": j.get("company_name", ""), **_job(j.get("title"),
             j.get("candidate_required_location") or "", "Remote", j.get("url"),
             j.get("publication_date")), "source": "Remotive"} for j in d["jobs"]]


def feed_wwr():
    txt = http_get("https://weworkremotely.com/remote-jobs.rss", as_json=False)
    if not txt:
        return []
    try:
        root = ET.fromstring(txt.encode("utf-8") if isinstance(txt, str) else txt)
    except ET.ParseError:
        return []
    out = []
    for it in root.iter("item"):
        title = it.findtext("title") or ""
        company, _, role = title.partition(":")
        if not role:
            company, role = "", title
        out.append({"company": company.strip(), **_job(role.strip(),
                    it.findtext("region") or "", "Remote", it.findtext("link"),
                    it.findtext("pubDate")), "source": "We Work Remotely"})
    return out


def feed_himalayas():
    d = http_get("https://himalayas.app/jobs/api?limit=100")
    if not d or "jobs" not in d:
        return []
    out = []
    for j in d["jobs"]:
        locs = j.get("locationRestrictions") or []
        locs = [l if isinstance(l, str) else (l.get("name") or "") for l in locs]
        out.append({"company": j.get("companyName", ""), **_job(
            j.get("title"), ", ".join(locs) or "Worldwide", "Remote",
            j.get("applicationLink") or j.get("guid"), j.get("pubDate")),
            "source": "Himalayas"})
    return out


FEEDS = [("RemoteOK", feed_remoteok), ("Remotive", feed_remotive),
         ("We Work Remotely", feed_wwr), ("Himalayas", feed_himalayas)]


def load_feeds(on_step=None):
    jobs = []
    for i, (name, fn) in enumerate(FEEDS, 1):
        if on_step:
            on_step(f"Loading remote job board: {name}")
        try:
            jobs += fn()
        except Exception:  # one broken feed must never stop the run
            pass
    return jobs

# ------------------------------------------------------- discussions --

_REDDIT_TOKEN = {"value": None, "expires": 0}


def _reddit_token():
    """App-only OAuth token. Free: create a 'script' app at reddit.com/prefs/apps and set
    REDDIT_CLIENT_ID / REDDIT_CLIENT_SECRET. Needed when running on cloud servers, which
    Reddit often blocks for anonymous requests."""
    cid, secret = os.environ.get("REDDIT_CLIENT_ID"), os.environ.get("REDDIT_CLIENT_SECRET")
    if not cid or not secret:
        return None
    if _REDDIT_TOKEN["value"] and time.time() < _REDDIT_TOKEN["expires"] - 60:
        return _REDDIT_TOKEN["value"]
    try:
        r = SESSION.post("https://www.reddit.com/api/v1/access_token",
                         auth=(cid, secret), data={"grant_type": "client_credentials"},
                         headers=HEADERS, timeout=TIMEOUT)
        d = r.json() if r.status_code == 200 else {}
    except (requests.RequestException, ValueError):
        d = {}
    if d.get("access_token"):
        _REDDIT_TOKEN.update(value=d["access_token"],
                             expires=time.time() + int(d.get("expires_in", 3600)))
        return d["access_token"]
    return None


def reddit_get(path):
    """path like '/r/forhire/new.json?limit=100'. Uses OAuth when configured."""
    token = _reddit_token()
    if token:
        p = path.replace(".json?", "?").replace(".json", "")
        return http_get("https://oauth.reddit.com" + p,
                        headers={"Authorization": f"bearer {token}"})
    return http_get("https://www.reddit.com" + path)


def reddit_posts(path):
    d = reddit_get(path)
    out = []
    for ch in ((d or {}).get("data") or {}).get("children") or []:
        p = ch.get("data") or {}
        out.append({"title": p.get("title", ""), "text": p.get("selftext", "") or "",
                    "url": "https://www.reddit.com" + p.get("permalink", ""),
                    "date": to_date(p.get("created_utc")), "sub": p.get("subreddit", ""),
                    "author": p.get("author", ""), "flair": p.get("link_flair_text") or ""})
    return out


def phrase_regex(words):
    words = [w.strip() for w in words if w and w.strip()]
    if not words:
        return None
    return re.compile(r"(?<![a-z0-9])(" + "|".join(
        re.escape(w).replace(r"\ ", r"\s+").replace("'", "['’]?") for w in words)
        + r")(?![a-z0-9])", re.I)


class Matchers:
    """Everything compiled once per run from the user's settings."""
    def __init__(self, cfg):
        self.cats = compile_categories(cfg)
        self.pain = {k: rx for k, v in (cfg.get("pain_signals") or DEFAULT_PAIN).items()
                     for rx in [phrase_regex(v)] if rx}
        self.small = phrase_regex(cfg.get("smallbiz_words") or DEFAULT_SMALLBIZ)

    def pains(self, text):
        hits = [(m.start(), n) for n, rx in self.pain.items()
                for m in [rx.search(text or "")] if m]
        return [n for _, n in sorted(hits)]

    def segment(self, text, sub=""):
        if sub.lower() in SMALLBIZ_SUBS or (self.small and self.small.search(text or "")):
            return SEG_SMALL
        return ""


def classify_post(title, text, sub, m):
    """-> (kind, category, segment) or None. Skips people offering services."""
    full = f"{title} {text}"
    if OFFER_RE.search(title) or OFFER_RE.search(text[:200]):
        return None
    if sub.lower() in ("forhire", "slavelabour") and not re.search(
            r"\[(hiring|task)\]|hiring", title, re.I):
        return None
    cats, pains = categorize(full, m.cats), m.pains(full)
    seg = m.segment(full, sub)
    if INTENT_RE.search(full) and (cats or pains):
        return "ask", (cats or pains)[0], seg
    if pains:
        return "struggle", pains[0], seg
    if cats:
        return "mention", cats[0], seg
    return None


def reddit_company(company, m, lookback):
    q = urllib.parse.quote(f'"{company}"')
    posts = reddit_posts(f"/search.json?q={q}&sort=new&t=year&limit=50")
    nre, out = name_regex(company), []
    for p in posts:
        full = f"{p['title']} {p['text']}"
        if not nre.search(full) or not in_window(p["date"], lookback):
            continue
        r = classify_post(p["title"], p["text"], p["sub"], m)
        if r:
            out.append(ev(r[0], r[1], f"Reddit r/{p['sub']}", p["title"], p["url"],
                          p["date"], f"r/{p['sub']}: {p['title']}", segment=r[2]))
    return out


def hn_search(query, since_ts, n=50):
    q = urllib.parse.quote(query)
    d = http_get(f"https://hn.algolia.com/api/v1/search_by_date?query={q}&tags=comment"
                 f"&numericFilters=created_at_i>{since_ts}&hitsPerPage={n}")
    out = []
    for h in (d or {}).get("hits") or []:
        text = strip_html(h.get("comment_text"))
        out.append({"text": text, "story": h.get("story_title") or "",
                    "author": h.get("author", ""), "date": to_date(h.get("created_at")),
                    "url": f"https://news.ycombinator.com/item?id={h.get('objectID')}"})
    return out


def hn_hiring_kind(first_line):
    if not REMOTE_WORD.search(first_line):
        return "job_onsite"
    if WORLDWIDE.search(first_line):
        return "job_eligible"
    if REGION_LOCK.search(first_line):
        return "job_locked"
    return "job_remote"


def hn_event(c, m, company_hint=None):
    first = c["text"][:220]
    cat = categorize(c["text"], m.cats)
    if not cat and not m.pains(c["text"]):
        return None, None
    if "who is hiring" in c["story"].lower():
        company = company_hint or first.split("|")[0].strip()[:60]
        if not cat:
            return None, None
        return company, ev(hn_hiring_kind(first), cat[0], "Hacker News: Who is hiring",
                           first[:120], c["url"], c["date"], first[:200],
                           role=(first.split("|")[1].strip() if "|" in first else ""))
    if "wants to be hired" in c["story"].lower() or OFFER_RE.search(first):
        return None, None
    r = classify_post("", c["text"], "", m)
    if not r:
        return None, None
    return company_hint or f"HN user {c['author']}", ev(
        r[0], r[1], "Hacker News", first[:120], c["url"], c["date"], first[:200],
        segment=r[2])


def hn_company(company, m, lookback):
    since = int((datetime.now() - timedelta(days=lookback)).timestamp())
    nre, out = name_regex(company), []
    for c in hn_search(f'"{company}"', since):
        if not nre.search(c["text"]):
            continue
        _, e = hn_event(c, m, company_hint=company)
        if e:
            out.append(e)
    return out

# --------------------------------------------------------------- news --

def news_events(company, cats, lookback):
    q = urllib.parse.quote(f'"{company}" when:{lookback}d')
    txt = http_get(f"https://news.google.com/rss/search?q={q}&hl=en-US&gl=US&ceid=US:en",
                   as_json=False)
    if not txt:
        return []
    try:
        root = ET.fromstring(txt.encode("utf-8") if isinstance(txt, str) else txt)
    except ET.ParseError:
        return []
    nre, best = name_regex(company), {}
    for it in list(root.iter("item"))[:40]:
        title = it.findtext("title") or ""
        if not nre.search(title):
            continue
        date = to_date(it.findtext("pubDate"))
        for kind, rx in NEWS_RULES:
            if not rx.search(title):
                continue
            cat = categorize(title, cats)
            if kind == "hiring_news" and not cat:
                continue  # generic hiring news only counts if it's about our areas
            e = ev(kind, cat[0] if cat else "Growth", "News", title, it.findtext("link"),
                   date, title)
            if kind not in best or e["date"] > best[kind]["date"]:
                best[kind] = e
            break
    return list(best.values())

# --------------------------------------------------------- website / tech --

def site_events(company, cats):
    d = domain_of(company["website"])
    if not d:
        return []
    out, base = [], f"https://{d}"
    page = http_get(base, as_json=False) or ""
    low = page.lower()
    for tool, pats in TECH_PATTERNS.items():
        if any(p in low for p in pats):
            ghl = tool == "GoHighLevel"
            out.append(ev("tech_ghl" if ghl else "tech",
                          "GoHighLevel" if ghl else "CRM Automation", "Website",
                          f"Website runs {tool}", base, "", f"{d} uses {tool}", tool=tool))
    for builder, pats in SMALL_SITE_PATTERNS.items():
        if any(p in low for p in pats):
            out.append(ev("small_site", "Growth", "Website", f"Site built on {builder}",
                          base, "", f"{d} is built on {builder} (typical of small businesses)",
                          segment=SEG_SMALL))
            break
    for path in ("/careers", "/jobs"):
        txt = http_get(base + path, as_json=False)
        if not txt:
            continue
        text = strip_html(re.sub(r"(?is)<(script|style).*?</\1>", " ", txt))
        cat = categorize(text, cats)
        if cat and re.search(r"(hiring|join (us|our team)|open (roles|positions)|apply)",
                             text, re.I):
            out.append(ev("careers", cat[0], "Website", f"Careers page mentions {cat[0]}",
                          base + path, "", f"{d}{path} mentions {', '.join(cat[:3])}"))
            break
    return out

# ------------------------------------------------------------ analysis --

def analyse(name, events, lookback):
    seen, evs = set(), []
    for e in events:
        key = (e["kind"], e["url"] or e["title"])
        if key in seen or not in_window(e["date"], max(lookback, 90 if
                                                       e["kind"].startswith("job") else 0)):
            continue
        seen.add(key)
        e["points"] = max(1, round(KIND_POINTS[e["kind"]] * recency(e["date"])))
        e["label"] = KIND_LABEL[e["kind"]]
        evs.append(e)
    evs.sort(key=lambda e: (e["points"], e["date"]), reverse=True)

    by_cat = {}
    for e in evs:
        by_cat[e["category"]] = by_cat.get(e["category"], 0) + e["points"]
    score = sum(min(v, CATEGORY_CAP) for v in by_cat.values())
    segment = SEG_SMALL if any(e.get("segment") == SEG_SMALL for e in evs) else ""
    if segment and any(e["kind"] != "small_site" for e in evs):
        score += SMALLBIZ_BONUS  # small owners decide fast and need help most
    score = min(100, score)
    needs = [c for c, _ in sorted(by_cat.items(), key=lambda x: -x[1]) if c != "Growth"]

    kinds = {e["kind"] for e in evs}
    if "job_eligible" in kinds:
        remote_ev, track = "Hires remote worldwide/Asia", "Apply (Job)"
    elif "job_remote" in kinds:
        remote_ev, track = "Remote roles (verify location)", "Apply (Job) - confirm location"
    elif "job_locked" in kinds:
        remote_ev, track = "Remote but region-locked", "Pitch VA / Contractor"
    elif "job_onsite" in kinds:
        remote_ev, track = "Hiring, but not remote", "Pitch VA / Contractor"
    else:
        remote_ev, track = "No job post - direct need", "Pitch VA / Contractor"
    if not evs:
        track = "Park - recheck later"
    role_ev = next((e for e in evs if e["kind"] in ("job_eligible", "job_remote")
                    and e.get("role")), None) or next(
        (e for e in evs if e.get("role")), None)
    dates = [e["date"] for e in evs if e["date"]]
    return {"signals": evs, "score": score,
            "tier": "HOT" if score >= 50 else "WARM" if score >= 25 else "COLD",
            "track": track, "remote_evidence": remote_ev, "needs": needs,
            "segment": segment,
            "intent_count": sum(e["kind"] in ("ask", "struggle", "mention", "hiring_news")
                                or e["kind"].startswith("job") for e in evs),
            "newest": max(dates) if dates else "",
            "top_role": role_ev.get("role") if role_ev else ""}


def scan_company(c, cfg, m, feed_jobs):
    cats = m.cats
    lb = cfg["lookback_days"]
    board, slug, jobs = find_jobs(c)
    events = job_events(jobs, cats, board or "", lb)
    matched = [j for j in feed_jobs if same_company(j["company"], c["company"])]
    for src in {j["source"] for j in matched}:
        events += job_events([j for j in matched if j["source"] == src], cats, src, lb)
    events += reddit_company(c["company"], m, lb)
    events += hn_company(c["company"], m, lb)
    events += news_events(c["company"], cats, lb)
    events += site_events(c, cats)
    res = analyse(c["company"], events, lb)
    sources = sorted({e["source"] for e in res["signals"]})
    res.update(company=c, board=f"{board} ({slug})" if board else "",
               sources=", ".join(sources))
    return res


def run(companies, cfg, on_progress=None, should_stop=lambda: False):
    """Scan the user's own company list."""
    m = Matchers(cfg)
    step = lambda msg: on_progress and on_progress(0, len(companies), msg, None)
    feed_jobs = load_feeds(step)
    results = []
    for i, c in enumerate(companies, 1):
        if should_stop():
            break
        print(f"[{i}/{len(companies)}] {c['company']} ...", end=" ", flush=True)
        res = scan_company(c, cfg, m, feed_jobs)
        results.append(res)
        print(f"{res['tier']} ({res['score']}) - {', '.join(res['needs'][:2]) or '-'}")
        if on_progress:
            on_progress(i, len(companies), c["company"], res)
        time.sleep(cfg.get("delay_seconds", 1))
    return results


def discover(cfg, on_progress=None, should_stop=lambda: False):
    """No list needed: find companies, founders and small-business owners who recently
    asked for these services or described a problem these services fix."""
    m = Matchers(cfg)
    cats = m.cats
    lb = cfg["lookback_days"]
    since = int((datetime.now() - timedelta(days=lb)).timestamp())
    subs, kws = cfg["subreddits"], cfg["hn_keywords"]
    total = len(FEEDS) + len(subs) + len(kws)
    done = [0]
    groups = {}

    def step(label):
        done[0] += 1
        if on_progress:
            on_progress(done[0], total, label, None)

    def add(company, e, website="", contact=""):
        key = norm(company) or company.lower()
        g = groups.setdefault(key, {"company": {"company": company, "website": website,
                                                "ats_slug": "", "contact_name": "",
                                                "contact_email": "", "link": contact},
                                    "events": []})
        g["events"].append(e)

    # 1. Remote job boards
    for name, fn in FEEDS:
        if should_stop():
            break
        step(f"Remote job board: {name}")
        try:
            jobs = fn()
        except Exception:
            jobs = []
        for j in jobs:
            if j["company"]:
                for e in job_events([j], cats, j["source"], lb):
                    add(j["company"], e)
        time.sleep(cfg.get("delay_seconds", 1))

    # 2. Reddit: people asking for help
    for sub in subs:
        if should_stop():
            break
        step(f"Reddit r/{sub}")
        for p in reddit_posts(f"/r/{sub}/new.json?limit=100"):
            if not in_window(p["date"], lb):
                continue
            r = classify_post(p["title"], p["text"], p["sub"] or sub, m)
            if not r or r[0] not in ("ask", "struggle"):
                continue
            s_ = p["sub"] or sub
            add(f"Reddit r/{s_} - u/{p['author']}",
                ev(r[0], r[1], f"Reddit r/{s_}", p["title"], p["url"], p["date"],
                   f"r/{s_}: {p['title']}", segment=r[2], post=p["text"][:600]),
                contact=p["url"])
        time.sleep(max(2, cfg.get("delay_seconds", 1)))  # be polite to Reddit

    # 3. Hacker News: Who is hiring + asks
    for kw in kws:
        if should_stop():
            break
        step(f"Hacker News: {kw}")
        for c in hn_search(kw, since):
            company, e = hn_event(c, m)
            if e and (e["kind"] != "mention"):
                add(company, e, contact=c["url"])
        time.sleep(cfg.get("delay_seconds", 1))

    results = []
    for g in groups.values():
        res = analyse(g["company"]["company"], g["events"], lb)
        if not res["signals"]:
            continue
        res.update(company=g["company"], board="",
                   sources=", ".join(sorted({e["source"] for e in res["signals"]})))
        results.append(res)
    results.sort(key=lambda r: (r["score"], r["newest"]), reverse=True)
    return results[: cfg.get("discover_limit", 75)]

# ------------------------------------------------------------------ emails --

NEED_PHRASE = {
    "GoHighLevel": "GoHighLevel setup and automation",
    "CRM Automation": "CRM and sales automation",
    "Sales & Business Development": "sales and business development support",
    "Customer Success": "customer success and support",
    "Virtual / Admin Assistant": "virtual assistance, admin and data entry",
}
QUICK_WINS = {
    "GoHighLevel": ["Funnels, pipelines and workflows built properly in GoHighLevel",
                    "Missed-call text-back, booking and follow-up automations",
                    "Snapshot setup, sub-account cleanup and reporting"],
    "CRM Automation": ["Automated lead capture, routing and follow-up sequences",
                       "Clean pipeline stages with dashboards leadership actually uses",
                       "Integrations between your CRM, forms, email and calendar"],
    "Sales & Business Development": ["Lead research and targeted list building",
                                     "Outreach sequences and appointment setting",
                                     "Pipeline hygiene and weekly sales reporting"],
    "Customer Success": ["Onboarding and check-in workflows",
                         "Inbox and ticket triage with clear SLAs",
                         "Renewal and churn-risk alerts inside your CRM"],
    "Virtual / Admin Assistant": ["Inbox, calendar and day-to-day admin support",
                                  "Accurate CRM data entry and database cleanup",
                                  "SOPs and simple automations for repetitive tasks"],
}


SMALL_WINS = {
    "GoHighLevel": ["Your GHL account set up properly: pipeline, calendar, forms",
                    "Missed-call text-back and automatic follow-ups so no lead is lost",
                    "Simple dashboards so you can see what's working"],
    "CRM Automation": ["One simple place to track every lead and customer",
                       "Automatic follow-ups so nobody slips through the cracks",
                       "Your forms, email and calendar connected, with less manual work"],
    "Sales & Business Development": ["A list of the right prospects for your business",
                                     "Outreach and follow-up done for you, consistently",
                                     "Booked calls on your calendar, not just 'leads'"],
    "Customer Success": ["Fast, friendly replies to customer messages",
                         "Simple onboarding and check-in emails that run themselves",
                         "Review and feedback requests that bring customers back"],
    "Virtual / Admin Assistant": ["Inbox, calendar and admin handled for you",
                                  "Data entry and record-keeping done accurately",
                                  "Repetitive tasks turned into simple routines or automations"],
}
PAIN_SUBJECT = {
    "Virtual / Admin Assistant": "Taking the admin off your plate",
    "Sales & Business Development": "Helping you land more clients",
    "CRM Automation": "Stopping leads from slipping through the cracks",
    "Customer Success": "Keeping your customers happy (and coming back)",
}
PAIN_SHORT = {
    "Virtual / Admin Assistant": "a heavy admin load",
    "Sales & Business Development": "the challenge of finding new clients",
    "CRM Automation": "leads slipping through the cracks",
    "Customer Success": "keeping customers happy and coming back",
}
PAIN_OPENER = {
    "Virtual / Admin Assistant": ("Running a business while handling every email, invoice "
                                  "and spreadsheet yourself is exhausting, and it's the "
                                  "easiest part to hand off."),
    "Sales & Business Development": ("Getting steady clients is the hardest part of a "
                                     "growing business, and it usually comes down to "
                                     "consistent outreach and follow-up, which I can run "
                                     "for you."),
    "CRM Automation": ("Leads going cold because follow-ups get missed is one of the most "
                       "fixable problems in a small business, and usually a few simple "
                       "automations solve it."),
    "Customer Success": ("When customers wait too long or feel forgotten, they leave. A "
                         "simple support and check-in routine fixes most of that."),
}


def first_name(full):
    return full.split()[0] if full else ""


def build_email(cfg, company, result):
    name = company["company"]
    is_person = name.startswith(("Reddit ", "HN user "))
    display = "your team" if is_person else name
    hello = (f"Hi {first_name(company['contact_name'])}," if company.get("contact_name")
             else "Hi there," if is_person else "Hi [First name],")
    sig = (f"\n\nBest regards,\n{cfg['your_name']}\n{cfg['headline']}\n{cfg['email']}"
           + (f"\n{cfg['linkedin']}" if cfg.get("linkedin") else "")
           + (f"\n{cfg['portfolio']}" if cfg.get("portfolio") else ""))
    need = result["needs"][0] if result["needs"] else "CRM Automation"
    phrase = NEED_PHRASE.get(need, need.lower())
    top = result["signals"][0] if result["signals"] else None
    kind = top["kind"] if top else ""
    small = result.get("segment") == SEG_SMALL
    if small:
        creds = (f"I help small businesses and new founders with {phrase}. I bring "
                 f"{cfg['years']} years in sales and hands-on experience with {cfg['tools']}"
                 f", so things get set up properly without agency prices. "
                 f"{cfg['proof_point']}")
        wins = "\n".join("- " + w for w in SMALL_WINS.get(need, SMALL_WINS["CRM Automation"]))
        cta = ("I work flexibly: a few hours a week or per project, with no long contracts. "
               "Happy to start with one small task so you can see the result first. Want me "
               "to send over 2-3 quick ideas for your situation?")
    else:
        creds = (f"I'm a remote CRM & Sales Systems specialist with {cfg['years']} years in "
                 f"B2B sales and hands-on builds in {cfg['tools']}. {cfg['proof_point']}")
        wins = "\n".join("- " + w for w in QUICK_WINS.get(need,
                                                         QUICK_WINS["CRM Automation"]))
        cta = (f"Would a 15-minute call this week be useful? I can share 2-3 ideas specific "
               f"to {display}, no strings attached.")

    if result["track"].startswith("Apply (Job)") and result["top_role"]:
        role = result["top_role"]
        subject = f"{role} - {cfg['years']} yrs B2B sales + CRM automation"
        body = (f"{hello}\n\nI saw {name} is hiring a remote {role} and wanted to reach "
                f"out directly, as well as applying through the posting.\n\n{creds}\n\n"
                f"I work fully remote from Pakistan (GMT+5) and overlap comfortably with "
                f"EU and US-East hours.\n\nWould you be open to a 15-minute call this week? "
                f"I'm happy to walk you through what I'd build for {name} in my first 30 "
                f"days.")
        return subject, body + sig

    if kind == "ask":
        if is_person:
            subject = f"Re: your post - {phrase}"
            opener = (f"I saw your recent post (\"{top['title'][:90]}\") and this is "
                      f"exactly the kind of work I do.")
        else:
            subject = f"{phrase[0].upper() + phrase[1:]} for {name}"
            opener = (f"I noticed {name} has recently been looking for help with {phrase}. "
                      f"That's exactly the kind of work I do.")
    elif kind == "struggle":
        if is_person:
            sub = top["source"].replace("Reddit ", "") if top["source"].startswith(
                "Reddit") else "the community"
            where = f"I came across your recent post in {sub}."
        else:
            where = (f"I noticed {name} has been dealing with "
                     f"{PAIN_SHORT.get(need, 'this')} recently.")
        subject = PAIN_SUBJECT.get(need, f"Help with {phrase}")
        opener = where + " " + PAIN_OPENER.get(
            need, "This is exactly the kind of problem I help with.")
    elif kind in ("job_locked", "job_onsite"):
        role = top.get("role") or "this role"
        subject = f"Your {role} opening - a remote contractor option"
        opener = (f"I saw {name} is hiring for {role}. If it helps while you hire, or "
                  f"instead of a full-time hire, I can take on that work remotely as a "
                  f"contractor.")
    elif kind == "tech_ghl":
        subject = f"Getting more out of GoHighLevel at {name}"
        opener = (f"I noticed {name} runs on GoHighLevel. Most teams use a fraction of what "
                  f"it can do, and a few workflows usually save hours every week.")
    elif kind == "tech":
        tool = top.get("tool", "your CRM")
        subject = f"Getting more out of {tool} at {name}"
        opener = (f"I noticed {name} uses {tool}. I help teams turn it from a contact list "
                  f"into an automated sales engine.")
    elif kind in ("crm_news",):
        subject = f"Helping {name} get the most from the CRM rollout"
        opener = (f"I came across the recent news (\"{top['title'][:90]}\"). Rollouts are "
                  f"where pipelines and automations either get set up right or stay messy "
                  f"for years.")
    elif kind == "funding":
        subject = f"Congrats on the raise - scaling {name}'s sales engine"
        opener = (f"Congratulations on the recent funding news. The next push is usually "
                  f"scaling sales, and that's where {phrase} saves the most time.")
    elif kind in ("hiring_news", "careers", "mention"):
        subject = f"{phrase[0].upper() + phrase[1:]} for {name}"
        opener = (f"I noticed {name} has been focusing on {phrase} recently, so I wanted to "
                  f"offer a hand.")
    elif kind == "leader":
        subject = f"Supporting {name}'s new revenue leadership"
        opener = (f"I saw {name} recently brought on new sales leadership. New leaders "
                  f"usually want pipeline visibility fast, which means CRM work.")
    else:
        subject = f"{phrase[0].upper() + phrase[1:]} for {name}"
        opener = f"I help growing teams like {name} with {phrase}."

    body = (f"{hello}\n\n{opener}\n\n{creds}\n\nWhere I can help right away:\n{wins}\n\n"
            f"{cta}")
    return subject, body + sig

# ------------------------------------------------------------------ output --

HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(bold=True, color="FFFFFF")
TIER_FILL = {"HOT": "F8CBAD", "WARM": "FFE699", "COLD": "DDEBF7"}


def style_sheet(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def link(cell, url):
    if url:
        cell.hyperlink = url
        cell.font = Font(color="0563C1", underline="single")


LEAD_COLS = ["Company / lead", "Tier", "Score", "Track", "Needs", "Lead type",
             "Intent signals",
             "Newest signal", "Remote evidence", "Top signal", "Top signal link",
             "New this run?", "Sources", "Contact name", "Contact email",
             "Find hiring manager", "Status", "Sent date", "Follow-up date", "Notes"]


def write_outputs(cfg, results, label="List", seen=None, out_dir=None, emails=True):
    """Build the Excel tracker. If `seen` (dict) is given it is updated in place and no
    seen-file is touched (the web app stores it in the database)."""
    out_dir = out_dir or OUT_DIR
    email_dir = os.path.join(out_dir, "emails")
    os.makedirs(email_dir if emails else out_dir, exist_ok=True)
    external_seen = seen is not None
    if not external_seen:
        seen = {}
        if os.path.exists(SEEN_FILE):
            try:
                with open(SEEN_FILE, encoding="utf-8") as f:
                    seen = json.load(f)
            except (OSError, ValueError):
                seen = {}

    wb = Workbook()
    leads = wb.active
    leads.title = "Leads"
    leads.append(LEAD_COLS)
    sig_ws = wb.create_sheet("Signals")
    sig_ws.append(["Company / lead", "Signal", "Need", "Source", "Points", "Detail",
                   "Date", "Link", "New?"])
    em_ws = wb.create_sheet("Email Drafts")
    em_ws.append(["Company / lead", "Tier", "To", "Subject", "Body"])
    sent_col = get_column_letter(LEAD_COLS.index("Sent date") + 1)

    results.sort(key=lambda r: (r["score"], r.get("newest", "")), reverse=True)
    new_total = 0
    for r in results:
        c = r["company"]
        any_new = False
        for s in r["signals"]:
            key = f"{c['company'].lower()}|{s['kind']}|{s['url'] or s['detail'][:80]}"
            is_new = key not in seen
            if is_new:
                seen[key] = datetime.now().strftime("%Y-%m-%d")
                new_total += 1
                any_new = True
            s["new"] = is_new
            sig_ws.append([c["company"], s["label"], s["category"], s["source"],
                           s["points"], s["detail"], s["date"] or "-", s["url"],
                           "NEW" if is_new else ""])
            link(sig_ws.cell(sig_ws.max_row, 8), s["url"])
        r["any_new"] = any_new

        top = r["signals"][0] if r["signals"] else None
        is_person = c["company"].startswith(("Reddit ", "HN user "))
        li = "" if is_person else ("https://www.linkedin.com/search/results/people/?keywords="
                                   + urllib.parse.quote(f"{c['company']} " + cfg.get(
                                       "hiring_manager_titles", "Head of Sales")))
        leads.append([c["company"], r["tier"], r["score"], r["track"],
                      ", ".join(r["needs"]) or "-", r.get("segment") or "-",
                      r["intent_count"], r["newest"] or "-",
                      r["remote_evidence"],
                      f"{top['label']}: {top['detail']}" if top else "No signals found",
                      top["url"] if top else "", "NEW" if any_new else "", r["sources"],
                      c.get("contact_name", ""), c.get("contact_email", ""),
                      "Reply on the post" if is_person else "Search LinkedIn",
                      "To contact" if r["tier"] != "COLD" else "Parked", "", "", ""])
        row = leads.max_row
        leads.cell(row, 2).fill = PatternFill("solid", fgColor=TIER_FILL[r["tier"]])
        link(leads.cell(row, LEAD_COLS.index("Top signal link") + 1), top["url"] if top else "")
        link(leads.cell(row, LEAD_COLS.index("Find hiring manager") + 1), li or c.get("link", ""))
        fu = leads.cell(row, LEAD_COLS.index("Follow-up date") + 1)
        fu.value = f'=IF({sent_col}{row}="","",{sent_col}{row}+4)'
        fu.number_format = "yyyy-mm-dd"
        leads.cell(row, LEAD_COLS.index("Sent date") + 1).number_format = "yyyy-mm-dd"

        if r["tier"] != "COLD":
            subject, body = build_email(cfg, c, r)
            r["email"] = {"subject": subject, "body": body}
            em_ws.append([c["company"], r["tier"], c.get("contact_email") or (
                "Reply on the post" if is_person else "[add email]"), subject, body])
            em_ws.cell(em_ws.max_row, 5).alignment = Alignment(wrap_text=True,
                                                               vertical="top")
            if not emails:
                continue
            fname = re.sub(r"[^A-Za-z0-9_-]+", "_", c["company"])[:60] + ".txt"
            with open(os.path.join(email_dir, fname), "w", encoding="utf-8") as f:
                f.write(f"To: {c.get('contact_email') or '[add email]'}\n"
                        f"Subject: {subject}\n\n{body}\n")

    style_sheet(leads, [30, 8, 7, 24, 30, 22, 9, 12, 26, 60, 14, 9, 22, 18, 26, 18, 12, 12,
                        13, 30])
    style_sheet(sig_ws, [30, 28, 24, 22, 8, 80, 12, 14, 7])
    style_sheet(em_ws, [30, 8, 26, 50, 100])
    for row in leads.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (5, 6, 9, 10))

    stamp = datetime.now().strftime("%Y-%m-%d_%H%M")
    path = os.path.join(out_dir, f"Signal_Leads_{label}_{stamp}.xlsx")
    wb.save(path)
    if not external_seen:
        with open(SEEN_FILE, "w", encoding="utf-8") as f:
            json.dump(seen, f, indent=1)
    return path, new_total

# -------------------------------------------------------------------- main --

def main():
    cfg = load_config()
    if "--discover" in sys.argv:
        print(f"Signal Hunter - finding warm leads from the last {cfg['lookback_days']} "
              f"days\n")
        results = discover(cfg, on_progress=lambda d, t, label, r: print(f"[{d}/{t}] {label}"))
        label = "Discover"
    else:
        args = [a for a in sys.argv[1:] if not a.startswith("--")]
        src = args[0] if args else os.path.join(BASE_DIR, "companies.csv")
        if not os.path.exists(src):
            sys.exit(f"Company list not found: {src}")
        companies = load_companies(src)
        if not companies:
            sys.exit("No companies found in that file.")
        print(f"Signal Hunter - scanning {len(companies)} companies\n")
        results = run(companies, cfg)
        label = "List"
    path, new_total = write_outputs(cfg, results, label)
    hot = sum(r["tier"] == "HOT" for r in results)
    warm = sum(r["tier"] == "WARM" for r in results)
    print(f"\nDone. HOT: {hot}  WARM: {warm}  New signals this run: {new_total}")
    print(f"Tracker: {path}\nEmail drafts: {EMAIL_DIR}")


if __name__ == "__main__":
    main()
